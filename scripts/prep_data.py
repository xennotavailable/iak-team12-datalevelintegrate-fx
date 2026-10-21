import csv, json, glob
from openpyxl import load_workbook
import os
BASE=os.path.dirname(os.path.abspath(__file__))
U=os.path.join(BASE,'..','data_sumber')+os.sep
def rows_of(path):
    if path.endswith('.csv'):
        return list(csv.reader(open(path,encoding='utf-8-sig')))
    ws=load_workbook(path,read_only=True,data_only=True).active
    return [list(r) for r in ws.iter_rows(values_only=True)]
def num(x):
    try: return float(x)
    except: return None
def parse(path,keys):
    out={}
    for r in rows_of(path):
        if not r or not isinstance(r[0],str): continue
        v=[num(c) for c in r[1:1+len(keys)]]
        if r[0].strip().upper()==r[0].strip() and any(x is not None for x in v) and len(r[0])>3 and not r[0].startswith('38'):
            out[r[0].strip()]=dict(zip(keys,v))
    return out
src=[('Persentase_Penduduk_Berumur*.xlsx',['pi_u','pi_r','pi_t']),
     ('Persentase_Rumah_Tangga*.csv',['rt_u','rt_r','rt_t']),
     ('Persentase_Penduduk_yang_Memiliki*.csv',['hp_u','hp_r','hp_t']),
     ('4_4_1*.xlsx',['tik'])]
data={}
for g,k in src:
    p=glob.glob(U+g)[0]; d=parse(p,k)
    for n,v in d.items(): data.setdefault(n,{}).update(v)
nat=data.pop('INDONESIA')
names=list(data)
rows=[dict(n=n.title().replace('Dki','DKI').replace('Di ','DI ').replace('Kep.','Kep.'),**{k:v for k,v in data[n].items()}) for n in names]
open(os.path.join(BASE,'..','data.js'),'w',encoding='utf-8').write('const D='+json.dumps({'rows':rows,'nat':nat},separators=(',',':'))+';\n')
print('data.js dibuat:',len(rows),'provinsi')
